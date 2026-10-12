"""Generate Class Module Full Working Guide as a Word document (uploadable to Google Docs)."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUTPUT = os.path.join(os.path.dirname(__file__), "BIAW_Class_Module_Full_Working_Guide.docx")


def set_cell_shading(cell, hex_color):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color)
    shading.set(qn("w:val"), "clear")
    cell._tePr = cell._tc.get_or_add_tcPr()
    cell._tc.get_or_add_tcPr().append(shading)


def add_heading_styled(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    return h


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)
    run.font.name = "Calibri"
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            run.font.size = Pt(11)
            run.font.name = "Calibri"


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Number")
        for run in p.runs:
            run.font.size = Pt(11)
            run.font.name = "Calibri"


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(10)
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = val
            for p in cells[c_idx].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
    doc.add_paragraph()


def main():
    doc = Document()

    # Title
    title = doc.add_heading("BIAW Class Module — Full Working Guide", 0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    sub = doc.add_paragraph()
    r = sub.add_run("Step-by-step documentation of how the Class Registration system works end to end.")
    r.italic = True
    r.font.size = Pt(12)

    meta = doc.add_paragraph()
    meta.add_run("Systems: ").bold = True
    meta.add_run("Node.js / Express · Airtable · Stripe · Webflow · Microsoft Graph (support@biaw.com)")

    # 1
    add_heading_styled(doc, "1. What This System Does", 1)
    add_para(doc, "The Class Module is the backend that powers BIAW class listings, registrations, payments, waitlists, cancellations, and certificate workflows. It keeps Airtable (admin data), Stripe (payments), Webflow (public website), and email in sync.")
    add_bullets(doc, [
        "Create and publish in-person and online classes",
        "Register members, non-members, and ROII (free) participants",
        "Collect payments with Stripe Checkout (no sales tax currently)",
        "Manage waitlists and notify users when seats open",
        "Sync seats and class details to biaw.com (Webflow)",
        "Handle refunds, cancellations, admin bookings, and certificates",
    ])

    # 2
    add_heading_styled(doc, "2. Architecture Overview", 1)
    add_para(doc, "The app follows an MVC-style layout:")
    add_bullets(doc, [
        "Routes — define API endpoints",
        "Controllers — handle HTTP requests",
        "Services — Stripe, Webflow, Email, CESCL, Scheduler",
        "Models — Airtable access for Class, Payment, Member, Waitlist",
        "Config — database, Stripe, Webflow, email settings",
    ])
    add_para(doc, "Main data flow:")
    add_para(doc, "Airtable / Website / Stripe  →  Class Module Server  →  Airtable + Stripe + Webflow + Email", bold=True)

    # 3
    add_heading_styled(doc, "3. Server Startup", 1)
    add_numbered(doc, [
        "app.js loads environment variables and Express.",
        "Webhook routes (/webhook, /webhook/airtable-class) use raw body parsing for signature verification.",
        "Other routes use JSON body parsing and CORS.",
        "Scheduled tasks initialize (schedulerService).",
        "Server listens on PORT (default 8000).",
    ])

    # 4
    add_heading_styled(doc, "4. Class Creation / Publish (Admin)", 1)
    add_para(doc, "Endpoint: POST /webhook/airtable-class", bold=True)
    add_para(doc, "When an admin creates a class in Airtable and the webhook fires:")
    add_numbered(doc, [
        "Airtable sends class fields (Name, Field ID, prices, seats, Product Type, location, instructor, images, etc.).",
        "Server parses the webhook into a flat class object.",
        "SchedulerService.processClassFromWebhook skips if Member Price ID and Non-Member Price ID already exist.",
        "StripeService.createStripeProductsAndCoupon creates: 2 Stripe products, 2 prices, optional coupon/promo code, 2 payment links. (No tax codes.)",
        "WebflowService.addClassToWebflow creates 2 CMS items (Member + Non-Member pages), publishes them live, and stores price IDs, payment links, seats, URLs, and class type.",
        "Class.updateClassRecord writes Webflow item IDs, Stripe price IDs, coupon code, seats remaining, and Publish status back to Airtable.",
    ])
    add_para(doc, "Result: Class is live on biaw.com and ready for registration.", bold=True)

    # 5
    add_heading_styled(doc, "5. Class Update / Unpublish / Delete", 1)
    add_table(doc,
        ["Endpoint", "Purpose"],
        [
            ["POST /api/endpoint", "Sync Airtable field changes to matching Webflow class items"],
            ["POST /api/unpublish", "Mark Webflow items unpublished"],
            ["POST /api/delete", "Delete/unpublish Webflow items based on Publish status"],
            ["POST /api/special", "Special class handling"],
            ["POST /api/category-sync", "Sync categories from Airtable to Webflow"],
        ],
    )

    # 6
    add_heading_styled(doc, "6. Paid Class Registration", 1)
    add_para(doc, "Endpoint: POST /submit-class", bold=True)
    add_para(doc, "When a user submits the registration form on the website:")
    add_numbered(doc, [
        "Website sends name, email, member ID, priceId, seats, participant details (P1-Name…), class Field ID, etc.",
        "Server validates price ID, seats, and member ID.",
        "Looks up the class in Airtable by Field ID.",
        "Creates participant rows in Multiple Class Registration (Payment Status = Pending).",
        "Creates a Payment Record (Pending) linked to the class and seat records.",
        "Creates a Stripe Checkout session: price × seats, promo codes allowed, success → thank-you, cancel → payment-declined. Links session to payment record via client_reference_id. No sales tax.",
        "Returns checkoutUrl; user is redirected to Stripe to pay.",
    ])
    add_para(doc, "Important: Seats are NOT reduced until payment succeeds.", bold=True)

    # 7
    add_heading_styled(doc, "7. After Successful Payment (Stripe Webhook)", 1)
    add_para(doc, "Endpoint: POST /webhook  |  Event: checkout.session.completed", bold=True)
    add_numbered(doc, [
        "Verify Stripe webhook signature.",
        "Only process if metadata.class === 'classmodule'.",
        "Find payment record by client_reference_id.",
        "Update payment to Paid (amount, payment intent ID).",
        "Sync purchase record to Webflow Payment CMS.",
        "Process CESCL registration if applicable.",
        "Reduce class seats in Airtable.",
        "Update remaining seats on Webflow.",
        "Mark participant records as Paid.",
        "Send confirmation email to the user.",
        "Send admin alert email.",
    ])
    add_para(doc, "Result: Booking is complete.", bold=True)

    # 8
    add_heading_styled(doc, "8. ROII Free Registration", 1)
    add_para(doc, "Endpoint: POST /register-class", bold=True)
    add_numbered(doc, [
        "Validate ROII member ID.",
        "Create participant records (Payment Status = ROII Free).",
        "Check seats available; fail if not enough.",
        "Reduce seats immediately (no Stripe).",
        "Sync remaining seats to Webflow.",
        "Create payment record as ROII-Free.",
        "Sync to Webflow and secondary Airtable bases (base2 / base3).",
        "Send confirmation and admin alert emails.",
    ])

    # 9
    add_heading_styled(doc, "9. Waitlist", 1)
    add_para(doc, "Join — POST /waitlist", bold=True)
    add_numbered(doc, [
        "User joins when the class is full.",
        "Record saved in Joint waitlist table.",
        "Confirmation email sent.",
    ])
    add_para(doc, "When seats open (after cancel / refund / seat update):")
    add_numbered(doc, [
        "System finds waitlist entries for that class.",
        "Sends “seat available” notification emails.",
    ])

    # 10
    add_heading_styled(doc, "10. Cancel / Refund", 1)
    add_table(doc,
        ["Endpoint", "Use"],
        [
            ["POST /cancel-payment", "User/admin cancel; Stripe refund if paid"],
            ["POST /api/refund", "Admin refund"],
            ["POST /api/roiicancel", "Cancel ROII free booking"],
            ["POST /api/without", "Cancel without refund"],
        ],
    )
    add_para(doc, "Typical cancel steps:")
    add_numbered(doc, [
        "Find payment record.",
        "Refund in Stripe (if paid).",
        "Mark payment / participants cancelled.",
        "Restore seats in Airtable and Webflow.",
        "Notify waitlist if seats opened.",
        "Send cancellation emails.",
    ])

    # 11
    add_heading_styled(doc, "11. Admin Operations", 1)
    add_table(doc,
        ["Endpoint", "What it does"],
        [
            ["POST /api/trigger-bookings", "Process admin-created bookings (paid + ROII)"],
            ["POST /api/trigger-seats", "Apply seat add/reduce from Airtable status fields"],
            ["POST /api/trigger-coupons", "Create Stripe coupons for classes needing discounts"],
            ["POST /api/trigger-all", "Run bookings + seats + coupons together"],
        ],
    )

    # 12
    add_heading_styled(doc, "12. Certificates", 1)
    add_table(doc,
        ["Endpoint", "What it does"],
        [
            ["POST /api/certification", "Generate certificate"],
            ["POST /api/send-certificate-email", "Email certificate"],
            ["POST /api/merge-certification-cards", "Merge certification card images/data"],
        ],
    )

    # 13
    add_heading_styled(doc, "13. Where Data Lives", 1)
    add_table(doc,
        ["System", "Stores"],
        [
            ["Airtable", "Classes, members, payments, participants, waitlist, CESCL, categories"],
            ["Stripe", "Products, prices, payment links, checkout, refunds"],
            ["Webflow", "Public class pages + purchase records"],
            ["Email (Microsoft Graph)", "Confirmations, waitlist, reminders, admin alerts"],
        ],
    )
    add_para(doc, "Main Airtable tables:", bold=True)
    add_bullets(doc, [
        "Biaw Classes",
        "Payment Records",
        "Multiple Class Registration",
        "Members",
        "Joint waitlist",
        "CESCL class registered records",
        "Category",
    ])

    # 14
    add_heading_styled(doc, "14. End-to-End Life of One Paid Booking", 1)
    add_numbered(doc, [
        "Admin creates class in Airtable.",
        "Webhook → Stripe products/prices/links + Webflow pages.",
        "User opens class on biaw.com and submits the form.",
        "/submit-class → pending Airtable records + Stripe checkout URL.",
        "User pays on Stripe.",
        "Stripe webhook → mark Paid, reduce seats, emails, Webflow sync.",
        "Later: cancel/refund restores seats and notifies waitlist.",
    ])

    # 15
    add_heading_styled(doc, "15. Current Payment Behavior (Important)", 1)
    add_bullets(doc, [
        "No sales tax is calculated or charged (tax logic removed from the server).",
        "Checkout is a simple one-time payment: seat price × quantity, optional promo code.",
        "Payment links use billing_address_collection: auto (Stripe may optionally collect address).",
        "Online and in-person paid registrations both use the same no-tax checkout path.",
    ])

    # 16
    add_heading_styled(doc, "16. Full Route Map", 1)
    add_table(doc,
        ["Route", "Purpose"],
        [
            ["GET /", "Health / server ready"],
            ["POST /webhook/airtable-class", "Create/publish new class"],
            ["POST /submit-class", "Paid registration"],
            ["POST /register-class", "ROII free registration"],
            ["POST /webhook", "Stripe payment completed"],
            ["POST /waitlist", "Join waitlist"],
            ["POST /cancel-payment", "Cancel / refund"],
            ["POST /api/mail", "Payment reminder email"],
            ["POST /api/endpoint", "Sync class updates to Webflow"],
            ["POST /api/unpublish", "Unpublish class on Webflow"],
            ["POST /api/delete", "Delete class on Webflow"],
            ["POST /api/special", "Special class webhook"],
            ["POST /api/category-sync", "Category sync"],
            ["POST /api/trigger-*", "Admin batch jobs"],
            ["POST /api/certification*", "Certificates"],
            ["POST /api/payment-status-update", "Payment status webhook"],
        ],
    )

    # 17
    add_heading_styled(doc, "17. Key Source Files", 1)
    add_table(doc,
        ["File", "Role"],
        [
            ["app.js", "Server entry point"],
            ["routes/index.js", "Main API routes"],
            ["controllers/classController.js", "Paid + ROII registration"],
            ["controllers/webhookController.js", "Stripe + Airtable class create webhooks"],
            ["controllers/airtableWebhookController.js", "Update / unpublish / delete sync"],
            ["controllers/adminController.js", "Refunds, seats, coupons, admin bookings"],
            ["services/stripeService.js", "Stripe products, checkout, refunds"],
            ["services/webflowService.js", "Webflow CMS create/update/publish"],
            ["services/emailService.js", "All outbound emails"],
            ["services/schedulerService.js", "Class processing + waitlist helpers"],
            ["models/Class.js / Payment.js", "Airtable class & payment operations"],
        ],
    )

    footer = doc.add_paragraph()
    footer.add_run("\nDocument generated for BIAW Class Module. Upload this .docx to Google Drive → Open with Google Docs to get a native Google Document.").italic = True

    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()
